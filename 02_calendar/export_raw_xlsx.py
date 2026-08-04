#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Seed or refresh the working event-list Excel (Events List.xlsx) from a source file.

This is a thin wrapper around the copy step that extract_demo_data.py already does.
Use it when you only want to replace the working Excel without re-extracting data.
"""
import argparse
from pathlib import Path

from openpyxl import load_workbook


def parse_args():
    parser = argparse.ArgumentParser(description="Copy source Excel to Events List.xlsx")
    parser.add_argument(
        "--source",
        required=True,
        help="Path to the source Excel file.",
    )
    parser.add_argument(
        "--output-dir",
        default=str(Path(__file__).parent),
        help="Directory where Events List.xlsx lives.",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    source = Path(args.source)
    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "Events List.xlsx"
    if not source.exists():
        raise FileNotFoundError(f"Source not found: {source}")

    wb = load_workbook(source)
    if "Events List" not in wb.sheetnames:
        raise ValueError(f"'{source}' must contain a sheet named 'Events List'.")
    for sheet_name in wb.sheetnames[:]:
        if sheet_name != "Events List":
            del wb[sheet_name]
    wb.save(out)
    wb.close()
    print(f"Copied {source} → {out} ({out.stat().st_size / 1024 / 1024:.2f} MB)")
    print("  Kept only 'Events List'.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
