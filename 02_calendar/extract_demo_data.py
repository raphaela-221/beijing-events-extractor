#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Extract High-priority events from the canonical Events List.xlsx into JSON/JS.

The canonical source is `01_event_list_output/Events List.xlsx`. Step 2 reads
from it directly (no copy/overwrite) so the file remains the single source of
truth for human editing.

For one-off seeding from a different file, use:

    python3 extract_demo_data.py --source /path/to/other.xlsx

That will still copy the provided file into `Events List.xlsx` as a local mirror.

Usage:
    python3 extract_demo_data.py
    python3 extract_demo_data.py --source /path/to/raw.xlsx --output-dir ./out
"""
import argparse
import hashlib
import json
import shutil
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

# Shared dedup logic (leaf module, no heavy deps) so the calendar
# display layer and the extractor stay consistent.
sys.path.insert(0, str(Path(__file__).parent.parent / "01_event_extractor"))
from src import dedup_key
from dates_parser import parse_dates_field


def parse_args():
    parser = argparse.ArgumentParser(
        description="Extract High-priority events from the canonical Events List.xlsx for the calendar demo."
    )
    parser.add_argument(
        "--source",
        default=None,
        help="Path to a source Excel file. If provided, it is copied over Events List.xlsx first. "
             "If omitted, the canonical file (../01_event_list_output/Events List.xlsx) is used directly.",
    )
    parser.add_argument(
        "--output-dir",
        default=str(Path(__file__).parent),
        help="Directory for events_data.json, events_data.js and state/last_run.json.",
    )
    return parser.parse_args()


def canonical_source_path() -> Path:
    """Return the canonical manually-edited Events List Excel."""
    return Path(__file__).parent.parent / "01_event_list_output" / "Events List.xlsx"


def fmt_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v)


def duration_days(start, end):
    """Inclusive span in days between two YYYY-MM-DD strings; None if unparseable."""
    try:
        s = datetime.strptime(start, "%Y-%m-%d").date()
        e = datetime.strptime(end, "%Y-%m-%d").date()
        return (e - s).days + 1
    except (ValueError, TypeError):
        return None


def dedup_records(records):
    """Display-layer guard: drop duplicate records (all categories).

    Reuses the same key as the extractor and the Excel dedup (dedup_key), so
    all three layers agree. Vacation: same city + vacation-type + year = one
    event. Other: same headline + same dates = one event. Returns
    (kept_records, removed_count).
    """
    kept, removed = dedup_key.deduplicate_events(records)
    return kept, len(removed)


def extract_records(ws):
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    idx = {name: i for i, name in enumerate(header)}

    # Backward compatibility with older Step 1 outputs.
    for legacy, current in {"事件类型": "Topic", "备注/地点": "备注"}.items():
        if legacy in idx and current not in idx:
            idx[current] = idx[legacy]

    def cell(row, name):
        i = idx.get(name)
        return None if i is None or i >= len(row) else row[i]

    records, total = [], 0
    for row in rows:
        if all(c is None or c == "" for c in row):
            continue
        total += 1
        priority = str(cell(row, "Priority") or "").strip()
        if priority != "High":
            continue
        topic = str(cell(row, "Topic") or "")
        parts = topic.split("\n", 1)
        topic_zh = parts[0].strip() if parts else ""
        topic_en = parts[1].strip() if len(parts) > 1 else ""
        # 中小学春秋假 / 中小学寒暑假 合并为单一类别"中小学假期"（大 topic），
        # 原细分（春秋假 / 寒暑假）保留到 topic_subtype，仅用于事件级展示。
        # 颜色 / 图例 / 信号 / 泳道都按合并后的"中小学假期"计，避免两条同源条目。
        topic_subtype = ""
        if topic_zh in ("中小学春秋假", "中小学寒暑假"):
            topic_subtype = topic_zh.replace("中小学", "")  # 春秋假 / 寒暑假
            topic_zh = "中小学假期"
        start = fmt_date(cell(row, "Start Date"))
        end = fmt_date(cell(row, "End Date")) or start
        # 时长>20天的赛事/展览不进 High：跨月长周期事件不是离散"大事件"，
        # 在月历里会拉出超长条形（如海淀3x3 87天、科协年会 31天）。
        # 业务规则，仅作用于日历展示层；canonical Excel 的 Priority 字段不变。
        if topic_zh in ("体育赛事", "大型会议和展览"):
            span = duration_days(start, end)
            if span is not None and span > 20:
                continue
        # 离散场次日期（Dates 列）：8.14-16, 8.19, 8.21-23, 8.28-30 -> 日期数组。
        # 解析成功时用首末覆写 start/end（时间轴横条 / eventsInMonth 归类与月历一致）；
        # 为空则回退 start/end 连续区间模型（向后兼容，未填 Dates 的事件不受影响）。
        dates_raw = cell(row, "Dates")
        fy = fm = None
        if start:
            try:
                _parts = str(start).replace("/", "-").split("-")
                fy, fm = int(_parts[0]), int(_parts[1])
            except (ValueError, IndexError):
                pass
        dates_list = parse_dates_field(dates_raw, fy, fm)
        if dates_list:
            start, end = dates_list[0], dates_list[-1]
        rec = {
            "no": cell(row, "No."),
            "topic_zh": topic_zh,
            "topic_en": topic_en,
            "topic_subtype": topic_subtype,
            "start": start,
            "end": end,
            "headline": str(cell(row, "Headline") or ""),
            "keywords_zh": str(cell(row, "Event Keywords") or ""),
            "keywords_en": str(cell(row, "Event English Keywords") or ""),
            "description": str(cell(row, "Event Description") or ""),
            "link": str(cell(row, "Link") or ""),
        }
        if dates_list:
            rec["dates"] = dates_list
        records.append(rec)
    return records, total


def record_signature(record: dict) -> str:
    """Stable signature of a record for delta detection."""
    payload = json.dumps(
        {
            "topic_zh": record.get("topic_zh"),
            "topic_en": record.get("topic_en"),
            "topic_subtype": record.get("topic_subtype"),
            "start": record.get("start"),
            "end": record.get("end"),
            "headline": record.get("headline"),
            "keywords_zh": record.get("keywords_zh"),
            "keywords_en": record.get("keywords_en"),
            "description": record.get("description"),
            "link": record.get("link"),
            "dates": record.get("dates"),
        },
        ensure_ascii=False,
        sort_keys=True,
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]


def compute_delta(records: list, state_path: Path) -> tuple:
    """Compare current records with last run and return (state, changed_months)."""
    new_signatures = {}
    for r in records:
        sig = record_signature(r)
        month = r["start"][:7] if r.get("start") else "unknown"
        new_signatures[sig] = month

    old_state = {"version": 1, "row_signatures": {}, "month_themes": {}}
    if state_path.exists():
        try:
            old_state = json.loads(state_path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            pass

    old_signatures = set(old_state.get("row_signatures", {}).keys())
    new_set = set(new_signatures.keys())

    added = new_set - old_signatures
    removed = old_signatures - new_set

    changed_months = set()
    for sig in added:
        changed_months.add(new_signatures[sig])
    old_sig_to_month = old_state.get("row_signatures", {})
    for sig in removed:
        changed_months.add(old_sig_to_month.get(sig, "unknown"))

    # Preserve previously cached themes so incremental theme generation can reuse them.
    merged_state = {
        "version": 1,
        "row_signatures": new_signatures,
        "changed_months": sorted(changed_months),
        "month_themes": old_state.get("month_themes", {}),
    }
    return merged_state, sorted(changed_months)


def main():
    args = parse_args()
    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    state_dir = out_dir / "state"
    state_dir.mkdir(parents=True, exist_ok=True)
    state_path = state_dir / "last_run.json"

    if args.source:
        source = Path(args.source)
        if not source.exists():
            raise FileNotFoundError(f"Source Excel not found: {source}")
        working_excel = out_dir / "Events List.xlsx"
        shutil.copy2(source, working_excel)
        print(f"Copied source to working file: {working_excel}")

        # The working/output Excel only needs the editable event list. Drop any
        # other sheets (raw source tabs, COUNT, Keywords, etc.) so the shared file
        # stays small and reader-facing.
        wb_work = load_workbook(working_excel)
        if "Events List" not in wb_work.sheetnames:
            raise ValueError(
                f"'{working_excel}' must contain a sheet named 'Events List'."
            )
        for sheet_name in wb_work.sheetnames[:]:
            if sheet_name != "Events List":
                del wb_work[sheet_name]
        wb_work.save(working_excel)
        wb_work.close()
        print("Kept only 'Events List' in working Excel.")

        excel_to_read = working_excel
    else:
        source = canonical_source_path()
        if not source.exists():
            raise FileNotFoundError(
                f"Canonical Events List not found: {source}. "
                "Please create it in 01_event_list_output/ or pass --source."
            )
        excel_to_read = source
        print(f"Using canonical event list: {excel_to_read}")

    wb = load_workbook(excel_to_read, read_only=True, data_only=True)
    print("Sheets:", wb.sheetnames)

    ws = wb["Events List"]
    records, total = extract_records(ws)
    records, dup_removed = dedup_records(records)
    records.sort(key=lambda x: x["start"] or "9999")

    json_path = out_dir / "events_data.json"
    js_path = out_dir / "events_data.js"

    json_path.write_text(
        json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    js_path.write_text(
        "window.EVENTS = " + json.dumps(records, ensure_ascii=False) + ";\n",
        encoding="utf-8",
    )

    state, changed_months = compute_delta(records, state_path)
    source_stat = excel_to_read.stat()
    state["canonical_path"] = str(excel_to_read.resolve())
    state["canonical_mtime"] = source_stat.st_mtime
    state["canonical_size"] = source_stat.st_size
    state_path.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"\nTotal data rows: {total}")
    if dup_removed:
        print(f"Duplicates dropped (display guard): {dup_removed}")
    print(f"High records: {len(records)}")
    print(f"Wrote: {json_path}")
    print(f"Wrote: {js_path}")
    print(f"Wrote: {state_path}")
    print(f"Changed months: {', '.join(changed_months) if changed_months else '(none)'}")

    tc = Counter(r["topic_zh"] for r in records)
    print("\nTopics:")
    for t, c in tc.most_common():
        print(f"  {c:4d}  {t}")

    mc = Counter(r["start"][:7] for r in records if r["start"])
    print("\nMonthly High counts:")
    for m in sorted(mc):
        print(f"  {m}  {mc[m]}")

    # 断点处理回显：含 dates 字段的事件（月历按天精确显示，中间断点不标色块）
    dated = [r for r in records if r.get("dates")]
    if dated:
        print(f"\n含离散场次（Dates 列）的事件 {len(dated)} 条（月历按天精确显示，断点日不标色块）：")
        for r in dated:
            print(f"  [{r['start']}~{r['end']}] {r['headline'][:30]}  ({len(r['dates'])} 场)")

    starts = [r["start"] for r in records if r["start"]]
    if starts:
        print(f"\nDate range: {min(starts)} ~ {max(starts)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
