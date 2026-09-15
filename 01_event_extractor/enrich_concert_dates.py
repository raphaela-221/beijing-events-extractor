#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Enrich canonical Events List.xlsx：LLM 从 Event Description 抽离散场次日期写回 Dates 列。

安全写法（零损耗，不损坏 canonical）：
  - openpyxl 只读读 canonical（read_only=True，绝不 save）
  - LLM 抽 Dates -> parse_dates_field 校验 -> _compress 压成 "8.14-16, 8.19, ..."
  - 调 add_dates_column.write_dates() 直接改 XML 把 Dates 写到 R 列（表外）
  - 不 openpyxl save（避免削 sharedStrings/printerSettings/数据验证 + insert_cols 破坏 table）

自愈模式（默认）：已有 Dates 的行不再盲目跳过，而是校验其展开日期是否
都能在该行自己的描述里找到（月/日 匹配）。对不上的（历史行序错位 --
canonical 被 Step1 重建重排后，按旧行号写入的 R 值整体串位）自动清掉重抽。
  - 校验通过 / 描述无日期可判 -> 跳过（不花 LLM）
  - 校验失败 -> 清空该值，走 LLM 重抽
--reset 强制全量重抽（清空全部已有 R 值后逐行抽，用于错位大面积发生时）。

幂等：R 列不存在则首次写入时加表头（R1="Dates"）+ 写到 R 列。
只处理含演唱会/音乐节/演出/巡演 且 description 有日期线索 的行。
抽取结果用 parse_dates_field 校验，解析失败则跳过（不污染数据）。

用法：
  python3 enrich_concert_dates.py            # 自愈模式：只重抽校验失败的行
  python3 enrich_concert_dates.py --reset    # 全量重抽（清空全部 Dates 后重抽）
  python3 enrich_concert_dates.py --dry-run  # 只打印不写
  python3 enrich_concert_dates.py --excel /path/to/other.xlsx
"""
import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parent.parent  # 项目根
CANONICAL = ROOT / "01_event_list_output" / "Events List.xlsx"

# 先加载 .env（src 模块 import 时读 env var），再 import src
load_dotenv(dotenv_path=ROOT / ".env")

# dates_parser / add_dates_column 在 02_calendar/
sys.path.insert(0, str(ROOT / "02_calendar"))
from dates_parser import parse_dates_field  # noqa: E402
from add_dates_column import write_dates  # noqa: E402

# llm_client 在本目录 src/
sys.path.insert(0, str(Path(__file__).resolve().parent))
from src.llm_client import (  # noqa: E402
    call_llm,
    get_llm_model,
    print_usage_summary,
    call_qwen38,
)


SYSTEM_PROMPT = """你是演出日期抽取助手。从中文演出描述里抽取每一场演出的具体日期，输出 Dates 列格式。

输出格式（严格遵守）：
- 逗号分隔，每段都带月份：8.14-16, 8.19, 8.21-23, 8.28-30
- 连续段用 - 连首尾日：8.14-16 = 14/15/16 三天
- 单日：8.19
- 跨年月用完整明写：2026-08-30~2026-09-02
- 原描述是「2026/4/2」这类斜杠写法时，先转成「4.2」再输出
- 只输出日期段，不要解释、引号、代码块或多余文字
- 年份从描述取；同年同月每段都重复月份（便于回解析）

示例：
描述："演唱会将于2026年8月14日至16日、8月19日、8月21日至23日、8月28日至30日举办"
输出：8.14-16, 8.19, 8.21-23, 8.28-30

描述："演出定于2026年9月10日、11日、13日"
输出：9.10-11, 9.13

描述："2026年8月30日至9月2日连续"
输出：2026-08-30~2026-09-02
"""

DATE_HINT = re.compile(
    r"\d{1,2}月\d{1,2}日|\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}\.\d{1,2}|至|到"
)
CONCERT_HINT = re.compile(r"演唱会|音乐节|巡演|concert|演出")

# 描述日期三格式：8月14日 / 2026-08-14 / 2026/4/2（区间中间日不在描述文本里，由校验容差处理）
_DESC_DATE_RE = re.compile(
    r"(?:(\d{1,2})月(\d{1,2})日)"
    r"|(?:\d{4}-(\d{1,2})-(\d{1,2}))"
    r"|(?:\d{4}/(\d{1,2})/(\d{1,2}))"
)
# 「至25日」省略月份的区间尾（前文已有 X月Y日 锚定了月份）
_DESC_RANGE_TAIL_RE = re.compile(r"[至到~\-](\d{1,2})日")


def _desc_date_pairs(desc: str) -> set:
    """描述文本里的 (月, 日) 集合（两种写法都吃 + 至区间尾日）。

    "2024年5月18日至6月1日" -> {(5,18),(6,1)}（两个锚点都被主正则命中）
    "2月1日至25日" -> {(2,1),(2,25)}（尾日靠 RANGE_TAIL 补，月份继承前锚点）
    """
    out = set()
    last_month = None
    pos = 0
    text = desc or ""
    while pos < len(text):
        m = _DESC_DATE_RE.search(text, pos)
        if not m:
            # 剩余文本里找「至N日」尾巴
            if last_month is not None:
                for t in _DESC_RANGE_TAIL_RE.finditer(text, pos):
                    out.add((last_month, int(t.group(1))))
            break
        if m.group(1):  # 8月14日
            mo, d = int(m.group(1)), int(m.group(2))
            out.add((mo, d))
            last_month = mo
            # 紧随其后的「至25日」（同月尾日）
            tail = _DESC_RANGE_TAIL_RE.match(text, m.end())
            if tail:
                out.add((mo, int(tail.group(1))))
        elif m.group(3):  # 2026-08-14
            out.add((int(m.group(3)), int(m.group(4))))
            last_month = int(m.group(3))
        else:  # 2026/4/2（斜杠格式）
            out.add((int(m.group(5)), int(m.group(6))))
            last_month = int(m.group(5))
        pos = m.end()
    return out


def _existing_dates_valid(dates_str: str, desc: str, fy, fm) -> bool:
    """校验已有 R 值是否「属于这一行」：展开后的每个日期的 (月,日)
    要么直接出现在描述文本里，要么落在描述日期构成的区间内。

    容差两级：
      1. 同月区间：日落在该月描述日期的 min~max 之间（8.14-16 的中间日 8/15，
         描述只写首尾"8月14日至16日"）
      2. 全局区间：日落在全部描述日期的最早~最晚之间（跨月连续区间
         "5月18日至6月1日" 的 5/19..5/31 两个月里各只有一个锚点，
         单月区间罩不住，需全局区间兜底）
    宽松方向是安全的：真正的错位值（如 6.27 出现在 8 月演唱会的行上）
    不会落进别的月份的日期区间。描述里完全抽不到日期 -> False（无法证明归属）。
    """
    expanded = parse_dates_field(dates_str, fy, fm)
    if not expanded:
        return False
    desc_pairs = _desc_date_pairs(desc)
    if not desc_pairs:
        return False  # 无日期可判 -> 交给调用方决定
    by_month = {}
    for mo, d in desc_pairs:
        by_month.setdefault(mo, []).append(d)
    # 全局区间（按 (月,日) 序数比较，跨月连续区间用）
    gmin = min(desc_pairs)
    gmax = max(desc_pairs)
    for iso in expanded:
        y, m, d = iso.split("-")
        m, d = int(m), int(d)
        if (m, d) in desc_pairs:
            continue
        days = by_month.get(m)
        if days and min(days) <= d <= max(days):
            continue  # 同月区间中间日
        if gmin <= (m, d) <= gmax:
            continue  # 跨月连续区间中间日
        return False
    return True


def _compress(dates_list):
    """['2026-08-14',...] -> '8.14-16, 8.19, 8.21-23, 8.28-30'（每段带月，可读+可回解析）。"""
    parsed = []
    for s in sorted(set(dates_list)):
        y, m, d = s.split("-")
        parsed.append((int(y), int(m), int(d)))
    groups, cur = [], None
    for y, m, d in parsed:
        if cur and cur[0] == y and cur[1] == m:
            cur[2].append(d)
        else:
            cur = (y, m, [d])
            groups.append(cur)
    parts = []
    for _y, m, days in groups:
        days.sort()
        segs, ss, prev = [], days[0], days[0]
        for d in days[1:]:
            if d == prev + 1:
                prev = d
                continue
            segs.append((ss, prev))
            ss = prev = d
        segs.append((ss, prev))
        for a, b in segs:
            parts.append(f"{m}.{a}" if a == b else f"{m}.{a}-{b}")
    return ", ".join(parts)


def _parse_llm_output(raw):
    s = (raw or "").strip().strip("`").strip('"').strip("'").strip()
    if s.startswith("```"):
        s = s.split("\n", 1)[-1].rsplit("```", 1)[0].strip()
    return s


def _extract_dates_for_desc(desc: str, provider: str = "qwen") -> str:
    """对单条描述抽日期段文本。provider="qwen" 走 Qwen 主力（失败切 call_llm
    兜底链）；provider="deepseek" 直接走 call_llm 兜底链。返回净化后的日期段文本。"""
    prompt = f"描述：{desc}\n输出："

    if provider == "qwen":
        try:
            raw = call_qwen38(
                prompt,
                system=SYSTEM_PROMPT,
                max_tokens=512,
                temperature=0.0,
            )
            return _parse_llm_output(raw)
        except Exception as e:
            print(f"    Qwen 主力失败，切兜底链：{e}")

    resp = call_llm(
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": prompt},
        ],
        model=get_llm_model(),
        temperature=0.0,
        max_completion_tokens=512,
    )
    return _parse_llm_output(resp.choices[0].message.content)


def enrich(excel_path, dry_run=False, reset=False, provider="qwen"):
    """openpyxl 只读抽 Dates，write_dates 写 R 列。不 openpyxl save。

    reset=False（自愈）：已有 R 值先校验归属，失败的重抽。
    reset=True：清空全部已有 R 值，全量重抽。
    provider="qwen"（默认）：Qwen3.8-27B（POMP 内网免费）主力抽，失败自动切
        call_llm 兜底链（DeepSeek → Ark → mlamp → Qwen）。
    provider="deepseek"：跳过 Qwen 主力，直接走 call_llm 兜底链。
    """
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    if "Events List" not in wb.sheetnames:
        wb.close()
        raise ValueError(f"{excel_path} 无 'Events List' sheet")
    ws = wb["Events List"]

    # 读表头
    header = []
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        header = [str(c).strip() if c is not None else "" for c in row]
        break
    col = {n: i for i, n in enumerate(header)}  # 0-based index into row tuple

    # Dates 列位置检查：只允许 R 列（或不存在）
    dates_tuple_idx = None
    if "Dates" in header:
        dates_tuple_idx = header.index("Dates")
        dates_letter = get_column_letter(dates_tuple_idx + 1)
        if dates_letter != "R":
            wb.close()
            raise ValueError(
                f"已有 Dates 列在 {dates_letter} 列，不是 R。本脚本只写 R 列，"
                f"不与旧 openpyxl insert_cols 版混用。请先恢复无 Dates 列的干净 canonical。"
            )

    row_dates = {1: "Dates"}  # R1 表头（幂等：已有则替换）
    # 清空写回集合：校验失败/reset 的已有值 -> 写 "" 覆盖旧错值（而非留着）
    enriched = 0
    skipped_ok = 0        # 已有值且归属校验通过 -> 保留
    skipped_no_judge = 0  # 已有值但描述无日期无法校验 -> 保留（宁可留着人工看）
    stale_cleared = 0     # 校验失败被清空重抽的
    skipped_no_concert = 0
    failed = 0
    r = 1  # 当前行号（min_row=2 起，循环内 +1）

    for row in ws.iter_rows(min_row=2, values_only=True):
        r += 1

        def cell(name):
            idx = col.get(name)
            return row[idx] if idx is not None and idx < len(row) else None

        topic = str(cell("Topic") or "")
        desc = str(cell("Event Description") or "")
        headline = str(cell("Headline") or "")
        if not (CONCERT_HINT.search(topic) or CONCERT_HINT.search(desc) or CONCERT_HINT.search(headline)):
            skipped_no_concert += 1
            continue
        if not DATE_HINT.search(desc):
            continue

        # fallback 年/月从 Start Date 取
        sd = cell("Start Date")
        start_str = sd.strftime("%Y-%m-%d") if isinstance(sd, datetime) else (str(sd) if sd else None)
        fy = fm = None
        if start_str:
            try:
                p = str(start_str).replace("/", "-").split("-")
                fy, fm = int(p[0]), int(p[1])
            except (ValueError, IndexError):
                pass

        # ---- 已有 R 值的自愈判定 ----
        existing = None
        if dates_tuple_idx is not None and dates_tuple_idx < len(row):
            existing = row[dates_tuple_idx]
            if existing is not None and str(existing).strip():
                existing = str(existing).strip()
                if reset:
                    stale_cleared += 1  # reset 模式：全清，下面走 LLM 重抽
                elif _desc_date_pairs(desc) and _existing_dates_valid(existing, desc, fy, fm):
                    skipped_ok += 1
                    continue  # 归属校验通过 -> 保留
                elif not _desc_date_pairs(desc):
                    skipped_no_judge += 1
                    continue  # 描述无日期无法校验 -> 保留（8.16 手填类）
                else:
                    print(f"  ⚠️ row {r}: 已有值归属校验失败，重抽（旧值 {existing!r}）")
                    stale_cleared += 1
            else:
                existing = None

        try:
            content = _extract_dates_for_desc(desc, provider=provider)
        except Exception as e:
            print(f"  ⚠️ LLM 调用失败 row {r}: {e}")
            failed += 1
            # 校验失败又抽不成的行：清掉错值比留着串位值好（回退 start/end 区间）
            if existing is not None and not dry_run:
                row_dates[r] = ""
            continue

        dates_list = parse_dates_field(content, fy, fm)
        if not dates_list:
            print(f"  ⚠️ 抽取为空 row {r}: desc={desc[:30]!r} -> LLM: {content[:60]!r}")
            failed += 1
            if existing is not None and not dry_run:
                row_dates[r] = ""
            continue

        dates_str = _compress(dates_list)
        row_dates[r] = dates_str
        print(f"  ✓ row {r}: {dates_str}   ({desc[:30]}...)")
        enriched += 1

    wb.close()  # read-only，不 save

    if not dry_run and (enriched > 0 or dates_tuple_idx is None or stale_cleared > 0):
        n = write_dates(excel_path, row_dates)
        print(f"\n已写回 {excel_path}（{n} 个单元格写到 R 列，含表头）")
    elif dry_run:
        print(f"\n[dry-run] 将写 {len(row_dates)} 个单元格（含表头）到 R 列")

    print(
        f"\n[Enrich] enriched={enriched}  "
        f"skipped(校验通过)={skipped_ok}  "
        f"skipped(无法校验,保留)={skipped_no_judge}  "
        f"cleared(清空重抽)={stale_cleared}  "
        f"skipped(非演唱会)={skipped_no_concert}  "
        f"failed={failed}  reset={reset}  dry_run={dry_run}"
    )
    print_usage_summary()
    return enriched


def main():
    ap = argparse.ArgumentParser(description="LLM 从 Event Description 抽离散场次日期写回 Dates 列（R 列，零损耗 XML）。默认自愈：已有值归属校验失败才重抽。")
    ap.add_argument("--excel", default=str(CANONICAL), help=f"Excel 路径（默认 canonical：{CANONICAL}）")
    ap.add_argument("--dry-run", action="store_true", help="只打印不写回")
    ap.add_argument("--reset", action="store_true", help="清空全部已有 Dates 值，全量重抽（错位大面积发生时用）")
    ap.add_argument("--provider", choices=["qwen", "deepseek"], default="qwen",
                    help="主力模型：qwen（Qwen3.8-27B POMP 免费，失败切兜底链）| deepseek（直接走 DeepSeek→Ark→mlamp→Qwen 兜底链）")
    args = ap.parse_args()

    p = Path(args.excel)
    if not p.exists():
        print(f"ERROR: {p} 不存在")
        return 1
    print(f"Enriching: {p}  (dry_run={args.dry_run}, reset={args.reset}, provider={args.provider})\n")
    enrich(p, dry_run=args.dry_run, reset=args.reset, provider=args.provider)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
