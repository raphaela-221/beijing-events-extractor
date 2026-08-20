"""canonical Events List.xlsx 核对往返：preview / download / upload / commit。

铁律（CLAUDE.md canonical-xlsx-openpyxl-gotcha）：
- canonical 含手工 Excel 特性（Table2/sharedStrings/printerSettings/数据验证）
- openpyxl load+save 会削 sharedStrings/printerSettings -> Excel 修复框
- 只读用 read_only=True，绝不 save
- 写回（commit）直接拷贝用户上传文件落盘，服务端绝不重新生成、绝不 save
- 结构校验跑 verify_canonical.py --excel <path> --quiet（0=完好 1=损坏 2=不存在）
"""
from __future__ import annotations

import hashlib
import secrets
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
from fastapi import APIRouter, Body, HTTPException, Query, Request
from fastapi.responses import FileResponse

from .. import config, db
from ..auth import require_user
from ..jobs.registry import list_jobs, resolve_upload_dir
from ..lock import lock as edit_lock
from ..runs.store import runs_since

router = APIRouter(prefix="/api/canonical", tags=["canonical"])

# canonical 列索引（A=0 起）
COL_NO, COL_TOPIC, COL_LINK, COL_START, COL_END = 0, 1, 2, 3, 4
COL_PRIORITY, COL_KEYWORDS, COL_DESC, COL_HEADLINE = 5, 6, 8, 9
COL_DATES = 17  # R

# diff 比对的字段：(列索引, 字段名, _read_all 的 key)
DIFF_FIELDS = [
    (COL_TOPIC, "Topic", "topic"),
    (COL_START, "Start Date", "start_date"),
    (COL_END, "End Date", "end_date"),
    (COL_PRIORITY, "Priority", "priority"),
    (COL_KEYWORDS, "Event Keywords", "keywords"),
    (COL_DESC, "Event Description", "desc"),
    (COL_HEADLINE, "Headline", "headline"),
    (COL_DATES, "Dates", "dates"),
]


# ---------- 只读取行 ----------
def _str(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _fmt_date(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _read_all(path: Path) -> list[dict]:
    """read_only 读全部事件行。绝不 save。"""
    wb = openpyxl.load_workbook(str(path), read_only=True)
    ws = wb.active
    out: list[dict] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or r[COL_NO] is None:
            continue
        out.append({
            "no": _str(r[COL_NO]),
            "topic": _str(r[COL_TOPIC]),
            "headline": _str(r[COL_HEADLINE]),
            "start_date": _fmt_date(r[COL_START]),
            "end_date": _fmt_date(r[COL_END]),
            "priority": _str(r[COL_PRIORITY]),
            "keywords": _str(r[COL_KEYWORDS]),
            "desc": _str(r[COL_DESC]),
            "dates": _str(r[COL_DATES]) if len(r) > COL_DATES else "",
        })
    wb.close()
    return out


# ---------- 工具 ----------
def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _verify(xlsx_path: Path) -> int:
    """跑 verify_canonical.py --excel <path> --quiet。0=完好 1=损坏 2=不存在。"""
    script = config.CALENDAR_DIR / "verify_canonical.py"
    if not xlsx_path.exists():
        return 2
    try:
        r = subprocess.run(
            [sys.executable, str(script), "--excel", str(xlsx_path), "--quiet"],
            capture_output=True, timeout=30,
        )
        return r.returncode
    except Exception:
        return 1


def _find_upload_file(upload_id: str) -> Path:
    """upload_id -> 上传目录里唯一的 xlsx 文件。"""
    d = resolve_upload_dir(upload_id)
    xlsx_files = sorted(d.glob("*.xlsx"))
    if not xlsx_files:
        raise HTTPException(400, "上传目录无 xlsx 文件")
    return xlsx_files[0]


def _record_download(hash_val: str, operator: str) -> str:
    download_id = secrets.token_hex(8)
    conn = db.connect()
    try:
        conn.execute(
            "INSERT INTO canonical_downloads (download_id, hash, downloaded_at, operator) VALUES (?,?,?,?)",
            (download_id, hash_val, datetime.now().isoformat(timespec="seconds"), operator),
        )
        conn.commit()
    finally:
        conn.close()
    return download_id


def _get_download(download_id: str) -> Optional[dict]:
    conn = db.connect()
    try:
        r = conn.execute(
            "SELECT download_id, hash, downloaded_at, operator FROM canonical_downloads WHERE download_id=?",
            (download_id,),
        ).fetchone()
        return dict(r) if r else None
    finally:
        conn.close()


def _backup() -> Path:
    """落盘前备份当前 canonical 到 CANONICAL_BACKUPS_DIR。"""
    config.ensure_dirs()
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    bak = config.CANONICAL_BACKUPS_DIR / f"Events List.{ts}.bak.xlsx"
    shutil.copy2(str(config.CANONICAL_PATH), str(bak))
    return bak


def _diff(canonical_path: Path, upload_path: Path) -> dict:
    """比对 canonical vs 上传文件，算 added/deleted/modified + 明细。绝不 save。"""
    cur = {r["no"]: r for r in _read_all(canonical_path)}
    new = {r["no"]: r for r in _read_all(upload_path)}
    added: list = []
    deleted: list = []
    modified: list = []
    for no, nr in new.items():
        if no not in cur:
            added.append(no)
        else:
            cr = cur[no]
            changes = []
            for _ci, fname, key in DIFF_FIELDS:
                cv = cr.get(key, "")
                nv = nr.get(key, "")
                if cv != nv:
                    changes.append({"field": fname, "old": cv, "new": nv})
            if changes:
                modified.append({"no": no, "headline": nr["headline"], "changes": changes})
    for no in cur:
        if no not in new:
            deleted.append(no)
    items = []
    for no in added:
        items.append({"type": "add", "no": no, "headline": new[no]["headline"], "field": "-", "change": "整行新增"})
    for no in deleted:
        items.append({"type": "del", "no": no, "headline": cur[no]["headline"], "field": "-", "change": "整行删除"})
    for m in modified:
        for c in m["changes"]:
            items.append({"type": "mod", "no": m["no"], "headline": m["headline"], "field": c["field"], "change": f"{c['old']} -> {c['new']}"})
    return {
        "added": len(added),
        "deleted": len(deleted),
        "modified": len(modified),
        "items": items,
    }


# ---------- 接口 ----------
@router.get("/preview")
def preview(
    request: Request,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    search: str = Query(""),
    priority: str = Query(""),
):
    """只读分页预览 canonical（绝不 save）。"""
    require_user(request)
    rows = _read_all(config.CANONICAL_PATH)
    if priority:
        p = priority.lower()
        rows = [r for r in rows if r["priority"].lower() == p]
    if search:
        s = search.lower()
        rows = [r for r in rows if s in (r["headline"] + r["topic"] + r["keywords"]).lower()]
    total = len(rows)
    start = (page - 1) * page_size
    items = rows[start:start + page_size]
    return {"items": items, "total": total, "page": page, "page_size": page_size}


@router.get("/row")
def get_row(request: Request, row: int = Query(..., ge=2)):
    """按 Excel 物理行号读单行（row_ref 回显用）。read_only，绝不 save。
    row=2 起为第一条数据（第 1 行是表头）。"""
    require_user(request)
    if not config.CANONICAL_PATH.exists():
        raise HTTPException(404, "canonical 文件不存在")
    wb = openpyxl.load_workbook(str(config.CANONICAL_PATH), read_only=True)
    try:
        ws = wb.active
        for r in ws.iter_rows(min_row=row, max_row=row, values_only=True):
            if not r or r[COL_NO] is None:
                return {"row": row, "exists": False}
            return {
                "row": row,
                "exists": True,
                "no": _str(r[COL_NO]),
                "topic": _str(r[COL_TOPIC]),
                "headline": _str(r[COL_HEADLINE]),
                "start_date": _fmt_date(r[COL_START]),
                "end_date": _fmt_date(r[COL_END]),
                "dates": _str(r[COL_DATES]) if len(r) > COL_DATES else "",
                "desc": _str(r[COL_DESC]) if len(r) > COL_DESC else "",
            }
        return {"row": row, "exists": False}
    finally:
        wb.close()


@router.get("/verify")
def verify(request: Request):
    """跑 verify_canonical.py --quiet。code: 0=完好 1=损坏 2=不存在。"""
    require_user(request)
    code = _verify(config.CANONICAL_PATH)
    detail = {0: "结构完好", 1: "canonical 结构损坏", 2: "文件不存在"}.get(code, "未知")
    return {"code": code, "detail": detail}


@router.get("/download")
def download(request: Request):
    """下载 canonical + 算 sha256 + 记 canonical_downloads 表。响应头带 hash + download_id。"""
    u = require_user(request)
    if not config.CANONICAL_PATH.exists():
        raise HTTPException(404, "canonical 文件不存在")
    h = _sha256(config.CANONICAL_PATH)
    download_id = _record_download(h, u["username"])
    resp = FileResponse(
        str(config.CANONICAL_PATH),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="Events List.xlsx",
    )
    resp.headers["X-Canonical-Hash"] = h
    resp.headers["X-Download-Id"] = download_id
    return resp


@router.post("/upload")
def upload(request: Request, body: dict = Body(...)):
    """上传回传：verify_canonical 校验 + 哈希比对（下载时 vs 当前）+ diff 摘要。绝不 save。"""
    require_user(request)
    upload_id = body.get("upload_id")
    download_id = body.get("download_id")
    if not upload_id:
        raise HTTPException(400, "缺少 upload_id")
    upath = _find_upload_file(upload_id)
    code = _verify(upath)
    if code != 0:
        detail = "结构损坏（openpyxl 往返或 Table2/sharedStrings 丢失）" if code == 1 else "文件不存在"
        raise HTTPException(400, f"上传文件 canonical 结构校验失败：{detail}")
    current_hash = _sha256(config.CANONICAL_PATH)
    downloaded_hash: Optional[str] = None
    downloaded_at: Optional[str] = None
    intervening: list = []
    if download_id:
        row = _get_download(download_id)
        if row:
            downloaded_hash = row["hash"]
            downloaded_at = row["downloaded_at"]
            if downloaded_hash != current_hash:
                write_ops = [j.id for j in list_jobs() if j.danger in ("write_canonical", "destructive")]
                intervening = runs_since(downloaded_at, write_ops)
    hash_match = downloaded_hash is None or downloaded_hash == current_hash
    diff = _diff(config.CANONICAL_PATH, upath)
    return {
        "upload_ok": True,
        "verified": True,
        "hash_match": hash_match,
        "current_hash": current_hash,
        "downloaded_hash": downloaded_hash,
        "downloaded_at": downloaded_at,
        "intervening_runs": intervening,
        "diff": diff,
    }


@router.post("/commit")
def commit(request: Request, body: dict = Body(...)):
    """写回 canonical：检查锁 -> .bak -> 落盘用户文件 -> 再 verify。绝不 openpyxl save。"""
    u = require_user(request)
    holder = edit_lock.status()
    if not holder or holder.get("user") != u["username"]:
        raise HTTPException(403, "需要先获取编辑锁才能写回 canonical")
    upload_id = body.get("upload_id")
    if not upload_id:
        raise HTTPException(400, "缺少 upload_id")
    upath = _find_upload_file(upload_id)
    backup_path = _backup()
    shutil.copy2(str(upath), str(config.CANONICAL_PATH))
    code = _verify(config.CANONICAL_PATH)
    if code != 0:
        shutil.copy2(str(backup_path), str(config.CANONICAL_PATH))
        raise HTTPException(500, "落盘后结构校验失败，已从 .bak 回滚")
    return {
        "committed": True,
        "backup_path": str(backup_path),
        "verified": True,
    }
